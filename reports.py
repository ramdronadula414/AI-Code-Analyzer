"""One complete report shared by text, paginated PDF and Word exports."""

import io
import json
import re
from xml.sax.saxutils import escape


def report_sections(result):
    meta, correction = result["metadata"], result["correction"]
    yield "AI Code Analyzer v2 — Security & Correction Report", [
        f"File: {meta['filename']}", f"Analyzed (UTC): {meta['analyzed_at']}",
        f"Source SHA-256: {meta['source_sha256']}",
        f"Scope: {meta['source_lines']} lines / {meta['source_characters']} characters / {meta['source_kind']}",
        f"Language: {result['language']}", f"Engine: {meta['engine']}", f"Model: {meta['model']}",
        f"Provider status: {meta['provider_status']} — {meta['provider_message']}",
        f"Duration: {meta['duration_seconds']} seconds", f"App version: {meta['version']}",
    ]
    yield "Executive summary", [f"Assessment: {result['classification']}",
                                f"Highest reported severity: {result['severity']}", result["summary"]]
    yield "Finding summary", [
        f"{f['id']} | {f['severity']} | {f['category']} | Lines {f['line_start']}-{f['line_end']} | {f['title']}"
        for f in result["findings"]
    ] or ["No findings reported. This does not establish that the code is safe."]
    for f in result["findings"]:
        yield f"{f['id']}: {f['title']}", [
            f"Category: {f['category']} | Severity: {f['severity']} | Confidence: {f['confidence']} (estimate)",
            f"Origin: {f['origin']} | Weakness: {f['cwe'] or 'Not specified'}",
            f"Location: lines {f['line_start']}-{f['line_end']} | Evidence: {f['evidence_status']}",
            f"Remediation: {f['remediation_status']}",
            "Source evidence:\n" + f["evidence"], "Explanation: " + f["explanation"],
            "Potential impact: " + f["impact"], "Recommended fix: " + f["recommendation"],
        ]
    if result["local_cross_check"]:
        yield "Local cross-check (indicators; not confirmed malware)", [
            f"{f['id']} | {f['severity']} | Line {f['line_start']} | {f['title']}\nEvidence: {f['evidence']}\nReview: {f['recommendation']}"
            for f in result["local_cross_check"]
        ]
    yield "Recommended actions", result["recommendations"] or ["Review the code and its trust boundaries; run application tests."]
    yield "Correction status", [f"Status: {correction['status']}", correction["reason"],
                                 "Generated code is a proposal. It has not been proven safe or functionally equivalent."]
    yield "Change log", [f"{', '.join(c['finding_ids'])}: {c['description']}" for c in correction["changes"]] or ["No accepted code changes."]
    yield "Validation performed", [
        f"Syntax check: {correction['validation']['status']} — {correction['validation']['detail']}",
        "Runtime tests: Not run. Neither original nor proposed code was executed.",
        (f"Local rescan: {len(correction['rescan_findings'])} remaining indicator(s). Absence of matches is not proof of safety."
         if correction["status"] == "proposed" else "Local rescan: Not run; no accepted replacement."),
    ]
    if correction["rescan_findings"]:
        yield "Remaining local indicators in proposed code", [
            f"{f['severity']} | Line {f['line_start']} | {f['title']}\nEvidence: {f['evidence']}\nReview: {f['recommendation']}"
            for f in correction["rescan_findings"]
        ]
    yield "Remaining risks and setup requirements", correction["remaining_risks"] or ["No additional risks supplied; independent review is still required."]
    yield "Suggested tests (not executed)", correction["suggested_tests"] or ["Test normal behavior, invalid inputs, permission boundaries and each identified issue in an isolated environment."]
    yield "Scope and limitations", result["limitations"]
    if correction["code"]:
        yield "Proposed corrected source (complete)", [correction["code"]]
        yield "Unified diff", [correction["diff"]]


def build_report_text(result):
    return "\n\n".join(title + "\n" + "=" * min(len(title), 80) + "\n" + "\n\n".join(items)
                       for title, items in report_sections(result))


def build_json_bytes(result):
    return json.dumps(result, ensure_ascii=False, indent=2).encode("utf-8")


def _xml_safe(text):
    # Word/PDF XML must not receive control characters from untrusted source/model text.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", text)


def build_pdf_bytes(result):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportBody", fontName="Helvetica", fontSize=9, leading=13,
                              alignment=TA_LEFT, splitLongWords=True, spaceAfter=5))
    styles.add(ParagraphStyle(name="ReportCode", fontName="Courier", fontSize=8, leading=11,
                              splitLongWords=True, spaceAfter=0))
    doc = SimpleDocTemplate(buffer, title="AI Code Analyzer v2 Security Report", author="AI Code Analyzer",
                            leftMargin=0.65 * inch, rightMargin=0.65 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.65 * inch)
    flow = []
    unicode_note = False
    def pdf_text(value):
        nonlocal unicode_note
        value = _xml_safe(value)
        # Built-in PDF fonts lack full Unicode. Make unsupported characters explicit instead of
        # silently dropping source characters. Exact UTF-8 code is in TXT/JSON/source exports.
        converted = []
        for char in value:
            try:
                char.encode("cp1252")
                converted.append(char)
            except UnicodeEncodeError:
                converted.append(f"[U+{ord(char):04X}]")
                unicode_note = True
        return escape("".join(converted)).replace("\t", "    ")
    for index, (title, items) in enumerate(report_sections(result)):
        flow.append(Paragraph(pdf_text(title), styles["Title" if index == 0 else "Heading2"]))
        code_section = title in ("Proposed corrected source (complete)", "Unified diff")
        for item in items:
            # One paragraph per bounded line allows arbitrary-length reports to paginate.
            for line in item.splitlines() or [""]:
                chunks = [line[i:i + 1000] for i in range(0, len(line), 1000)] or [""]
                for chunk in chunks:
                    rendered = pdf_text(chunk)
                    if code_section:
                        rendered = rendered.replace(" ", "&#160;")
                    flow.append(Paragraph(rendered or "&#160;", styles["ReportCode" if code_section else "ReportBody"]))
            flow.append(Spacer(1, 5))
    if unicode_note:
        flow.append(Paragraph("PDF font note: unsupported Unicode characters are shown as [U+XXXX]. Use the TXT, JSON, Word or source download for exact characters.", styles["ReportBody"]))
    def footer(canvas, document):
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.setFont("Helvetica", 8)
        canvas.drawString(document.leftMargin, 24, "AI Code Analyzer v2 | Static review | Human verification required")
        canvas.drawRightString(document.pagesize[0] - document.rightMargin, 24, f"Page {document.page}")
    doc.build(flow, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def build_docx_bytes(result):
    from docx import Document
    from docx.shared import Pt
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)
    for index, (title, items) in enumerate(report_sections(result)):
        document.add_heading(_xml_safe(title), 0 if index == 0 else 1)
        code_section = title in ("Proposed corrected source (complete)", "Unified diff")
        for item in items:
            paragraph = document.add_paragraph(_xml_safe(item))
            if code_section:
                for run in paragraph.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
