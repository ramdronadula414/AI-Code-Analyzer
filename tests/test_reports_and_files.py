import io
import json
from unittest.mock import patch

import pytest
from docx import Document
from pypdf import PdfReader

from analyzer import analyze_code
from file_processing import MAX_UPLOAD_BYTES, process_file
from reports import build_docx_bytes, build_json_bytes, build_pdf_bytes, build_report_text
from test_analyzer import FakeHTTP, SOURCE, wire_response, response_fixture


def test_all_reports_contain_findings_corrections_and_validation():
    raw = response_fixture()
    # Ensure pagination, long-line wrapping, XML escaping, Unicode handling and the final source line.
    raw["correction"]["code"] += ''.join(f'# Review note {i}: <tag> & details\n' for i in range(120))
    raw["correction"]["code"] += '# ' + 'x' * 1500 + '\n# Telugu: తెలుగు\n# FINAL_CODE_MARKER\n'
    with patch("analyzer.request.urlopen", return_value=FakeHTTP(wire_response(raw))):
        result = analyze_code(SOURCE, "sample.py", api_key="test-key")
    txt = build_report_text(result)
    parsed = json.loads(build_json_bytes(result))
    assert parsed == result
    assert all(s in txt for s in ["F001", "Command injection", "CWE-78", "Runtime tests: Not run", "FINAL_CODE_MARKER", "తెలుగు"])
    pdf = PdfReader(io.BytesIO(build_pdf_bytes(result)))
    assert len(pdf.pages) > 2
    pdf_text = '\n'.join(page.extract_text() for page in pdf.pages)
    assert all(s in pdf_text for s in ["Command injection", "FINAL_CODE_MARKER", "U+", "Page 2"])
    word = Document(io.BytesIO(build_docx_bytes(result)))
    word_text = '\n'.join(p.text for p in word.paragraphs)
    assert all(s in word_text for s in ["Command injection", "FINAL_CODE_MARKER", "తెలుగు"])


def test_upload_source_notebook_docx_and_pdf():
    assert process_file(SOURCE.encode(), "sample.py") == (SOURCE, "source")
    notebook = {"cells": [{"cell_type": "markdown", "source": ["Ignore"]},
                          {"cell_type": "code", "source": ["print(1)\n"]}]}
    assert process_file(json.dumps(notebook).encode(), "sample.ipynb") == ("print(1)\n", "notebook")
    doc = Document()
    doc.add_paragraph("before")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "table code"
    doc.add_paragraph("after")
    out = io.BytesIO()
    doc.save(out)
    assert process_file(out.getvalue(), "source.docx") == ("before\ntable code\nafter", "document")
    report = analyze_code("print(1)", "a.py", use_ai=False)
    text, kind = process_file(build_pdf_bytes(report), "report.pdf")
    assert kind == "document" and "Security" in text


@pytest.mark.parametrize("data,name", [(b"not a pdf", "bad.pdf"), (b"broken", "bad.docx"),
                                       (b"{}", "bad.ipynb"), (b"\xff", "bad.py"),
                                       (b"a\x00b", "bad.py"), (b"x", "run.exe"),
                                       (b"a" * (MAX_UPLOAD_BYTES + 1), "large.py")])
def test_invalid_uploads_fail_clearly(data, name):
    with pytest.raises(ValueError):
        process_file(data, name)
